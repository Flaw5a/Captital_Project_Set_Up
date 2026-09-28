"""Builders for tiny CBES-style template fixtures used by tests and the demo."""
from __future__ import annotations

from pathlib import Path

from docx import Document
import openpyxl
from pptx import Presentation
from pptx.util import Inches


def make_docx_form(path: Path, title="Site Safety & Env Mgt Plan") -> None:
    """A docx with a header block table mirroring the real CBES layout."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=1)
    t = doc.add_table(rows=3, cols=4)
    t.style = "Table Grid"
    t.cell(0, 0).text = "Project Title:"
    t.cell(0, 1).text = "OLD STALE SITE"          # stale value -> must be overwritten
    t.cell(0, 2).text = "Project Number:"
    t.cell(0, 3).text = ""
    t.cell(1, 0).text = "Site Location:"
    t.cell(1, 1).text = ""
    t.cell(1, 2).text = "Postcode:"
    t.cell(1, 3).text = ""
    t.cell(2, 0).text = "Project Manager:"
    t.cell(2, 1).text = ""
    t.cell(2, 2).text = "Quantity Surveyor:"
    t.cell(2, 3).text = ""
    doc.save(str(path))


def make_docx_simple(path: Path) -> None:
    """A docx with a single 'Project Name:' row (like MS 9.4)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "Project Name:"
    t.cell(0, 1).text = ""
    t.cell(1, 0).text = "Site:"
    t.cell(1, 1).text = ""
    doc.save(str(path))


def make_xlsx_form(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Project Number:"
    ws["B1"] = ""
    ws["A2"] = "Site:"
    ws["B2"] = ""
    ws["A3"] = "Project Manager:"
    ws["B3"] = ""
    wb.save(str(path))


def make_pptx_form(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "Project: TBC"
    prs.save(str(path))


def make_lookup_workbooks(root: Path) -> None:
    """Create demo PMs / QS / Customer lookup workbooks at the root."""
    root.mkdir(parents=True, exist_ok=True)
    pm = openpyxl.Workbook(); ws = pm.active; ws.title = "Project Manager"
    ws.append(["Full Name", "Employment ID"])
    for n, i in [("Andy Flaws", 200001), ("Brian Rooney", 200002), ("Shane Metcalf", 200003)]:
        ws.append([n, i])
    pm.save(str(root / "PMs.xlsx"))

    qs = openpyxl.Workbook(); ws = qs.active; ws.title = "QS"
    ws.append(["Full Name", "Employment ID"])
    for n, i in [("Alex Lee", 220001), ("Jordan Reid", 220002)]:
        ws.append([n, i])
    qs.save(str(root / "QS.xlsx"))

    cu = openpyxl.Workbook(); ws = cu.active; ws.title = "Customer List"
    ws.append(["Customer"])
    for c in ["Asda", "Co-op", "Marks & Spencer", "Morrisons", "Tesco"]:
        ws.append([c])
    cu.save(str(root / "CBES Customer List.xlsx"))


def build_sample_root(root: Path) -> None:
    """Create a minimal but realistic TEMPLATES_ROOT for the demo / e2e tests."""
    make_lookup_workbooks(root)
    master = root / "00 Master Projects FIle Template 0524 (Opt 2)"
    for dept in ("M&E", "M&ELite"):
        for sub in ("01 - Pre Construction", "02 - Construction Phase",
                    "03 - Handover", "04 - Health & Safety", "05 - Commercial"):
            (master / dept / sub).mkdir(parents=True, exist_ok=True)

    forms = root / "05 Forms"
    make_docx_form(forms / "01 OHS" / "HS 11.2 - Site Safety & Env Mgt Plan.docx")
    make_docx_simple(forms / "01 OHS" / "HS 11.1 Const Phase HSE Plan.docx")
    make_xlsx_form(forms / "02 ENV" / "ENV 1.2 - Waste Note.xlsx")
    make_docx_simple(forms / "03 QUAL" / "MS 9.4 Project Notification.docx")
    make_pptx_form(forms / "01 OHS" / "HS 13.1 - Site Fire Plan.pptx")
    make_docx_form(forms / "04 IMS" / "IMS 4.3 - PM Weekly Inspection.docx")
