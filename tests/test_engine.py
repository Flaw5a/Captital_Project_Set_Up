"""Engine tests: label matching, per-format stampers, and end-to-end generate()."""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pptx import Presentation

from app.engine.fields import FieldMatcher, ProjectData, normalise_label
from app.engine.stamp_docx import stamp_docx
from app.engine.stamp_pptx import stamp_pptx
from app.engine.stamp_xlsx import stamp_xlsx
from tests.fixtures import (
    build_sample_root,
    make_docx_form,
    make_pptx_form,
    make_xlsx_form,
)

FIELD_MAP = {
    "fields": [
        {"field": "project_number", "enabled": True, "labels": ["project number"]},
        {"field": "project_title", "enabled": True, "labels": ["project title", "project name", "project"]},
        {"field": "site", "enabled": True, "labels": ["site", "site location"]},
        {"field": "postcode", "enabled": True, "labels": ["postcode"]},
        {"field": "project_manager", "enabled": True, "labels": ["project manager"]},
        {"field": "quantity_surveyor", "enabled": True, "labels": ["quantity surveyor"]},
        {"field": "date", "enabled": False, "labels": ["date"]},
    ]
}

VALUES = {
    "project_number": "PJ-123456",
    "project_title": "Asda Stockton Refresh",
    "site": "Asda Stockton",
    "postcode": "TS18 2PB",
    "project_manager": "Andy Flaws",
    "quantity_surveyor": "Alex Lee",
    "client": "Asda",
    "date": "28/09/2026",
}


def matcher():
    return FieldMatcher(FIELD_MAP, VALUES, stamp_date=False)


# --------------------------------------------------------------------------- #
def test_normalise_label():
    assert normalise_label("  Project  Number: ") == "project number"
    assert normalise_label("PROJECT NO:") == "project no"
    assert normalise_label(" Site ") == "site"


def test_matcher_exact_only():
    m = matcher()
    assert m.value_for("Project Number:") == "PJ-123456"
    assert m.value_for("project number") == "PJ-123456"
    assert m.value_for("A random cell") is None
    # Empty project value -> no write (client not in field map here).
    assert m.value_for("Something") is None


def test_pj_validation():
    assert ProjectData("PJ-123456", "Site").validate() == []
    assert ProjectData("PJ-12", "Site").validate()          # too few digits
    assert ProjectData("123456", "Site").validate()         # missing prefix
    assert ProjectData("PJ-123456", "").validate()          # missing site


def test_project_title_defaults_to_site():
    d = ProjectData("PJ-000001", "Asda Bootle")
    assert d.project_title == "Asda Bootle"


# --------------------------------------------------------------------------- #
def test_stamp_docx_overwrites_and_fills(tmp_path):
    p = tmp_path / "form.docx"
    make_docx_form(p)
    n = stamp_docx(str(p), matcher())
    assert n >= 5
    doc = Document(str(p))
    t = doc.tables[0]
    assert t.cell(0, 1).text == "Asda Stockton Refresh"   # stale "OLD STALE SITE" replaced
    assert t.cell(0, 3).text == "PJ-123456"
    assert t.cell(1, 1).text == "Asda Stockton"           # Site Location
    assert t.cell(1, 3).text == "TS18 2PB"
    assert t.cell(2, 1).text == "Andy Flaws"
    assert t.cell(2, 3).text == "Alex Lee"


def test_stamp_xlsx(tmp_path):
    p = tmp_path / "form.xlsx"
    make_xlsx_form(p)
    n = stamp_xlsx(str(p), matcher())
    assert n >= 2
    wb = openpyxl.load_workbook(str(p))
    ws = wb.active
    assert ws["B1"].value == "PJ-123456"
    assert ws["B2"].value == "Asda Stockton"
    assert ws["B3"].value == "Andy Flaws"


def test_stamp_pptx_inline(tmp_path):
    p = tmp_path / "form.pptx"
    make_pptx_form(p)
    stamp_pptx(str(p), matcher())
    prs = Presentation(str(p))
    texts = [
        sh.text_frame.text
        for sl in prs.slides for sh in sl.shapes if sh.has_text_frame
    ]
    assert any("Asda Stockton Refresh" in t for t in texts)


def test_date_disabled_by_default(tmp_path):
    """A 'Date:' label must NOT be filled when stamp_date is False."""
    from tests.fixtures import make_docx_simple
    p = tmp_path / "d.docx"
    doc = Document()
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Date:"
    t.cell(0, 1).text = ""
    doc.save(str(p))
    stamp_docx(str(p), matcher())
    assert Document(str(p)).tables[0].cell(0, 1).text == ""


# --------------------------------------------------------------------------- #
def test_generate_end_to_end(tmp_path, monkeypatch):
    root = tmp_path / "root"
    out = tmp_path / "out"
    build_sample_root(root)

    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    monkeypatch.setenv("OUTPUT_ROOT", str(out))
    monkeypatch.setenv("STAMP_DATE", "false")

    import app.config as config
    config.get_settings.cache_clear()
    config.load_departments.cache_clear()
    config.load_field_map.cache_clear()

    from app.engine import generator, structure
    tree = structure.forms_tree()
    selected = [f["rel"] for grp in tree for f in grp["files"]]
    assert selected, "sample forms should be discovered"

    data = ProjectData(
        project_number="PJ-654321",
        site="Asda Stockton",
        postcode="TS18 2PB",
        project_manager="Andy Flaws",
        quantity_surveyor="Alex Lee",
        department="M&E",
        structure="standard",
    )
    report = generator.generate(data, selected)

    dest = Path(report.output_dir)
    assert dest.is_dir()
    assert dest.name == "PJ-654321 - Asda Stockton"
    # Department skeleton recreated.
    assert (dest / "04 - Health & Safety").is_dir()
    # Forms filed across the structure + filename prefixed.
    all_docx = list(dest.rglob("*.docx"))
    assert all_docx
    # Prefix is the project number ONLY (no site in the filename).
    assert all(f.name.startswith("PJ-654321 - ") for f in all_docx)
    assert all(" - Asda Stockton - " not in f.name for f in all_docx)
    assert report.total_fields > 0
    assert report.error_count == 0
    assert Path(report.zip_path).is_file()

    # OHS forms are now filed into Health & Safety, not the populated fallback.
    assert any("Site Safety" in f.name for f in (dest / "04 - Health & Safety").rglob("*.docx"))

    # Verify a stamped value survived into the filed copy.
    sample = next(f for f in all_docx if "Site Safety" in f.name)
    doc = Document(str(sample))
    vals = [c.text for tb in doc.tables for row in tb.rows for c in row.cells]
    assert "PJ-654321" in vals


def test_lookups(tmp_path, monkeypatch):
    root = tmp_path / "root"
    from tests.fixtures import make_lookup_workbooks
    make_lookup_workbooks(root)
    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    import app.config as config
    config.get_settings.cache_clear()
    config.load_departments.cache_clear()
    from app.engine import lookups
    lookups.clear_cache()
    pms = lookups.project_managers()
    assert "Andy Flaws" in pms and "Brian Rooney" in pms
    assert "Alex Lee" in lookups.quantity_surveyors()
    assert "Marks & Spencer" in lookups.customers()


def test_geo_links_and_graceful():
    from app.engine.geo import lookup_postcode
    # Invalid format -> not valid, but Google links are still built.
    bad = lookup_postcode("NOTAPOSTCODE")
    assert bad["valid"] is False
    assert "postcode" in bad["error"].lower()
    assert bad["google_earth_url"].startswith("https://earth.google.com/")
    # Well-formed postcode -> always returns clickable links (lookup may or may
    # not resolve depending on network egress; both paths are acceptable).
    ok = lookup_postcode("TS18 2PB")
    assert ok["google_maps_url"] and ok["google_earth_url"]
    assert ok["input"] == "TS18 2PB"


# --------------------------------------------------------------------------- #
def test_generate_files_by_group_into_subfolders(tmp_path, monkeypatch):
    """v2: forms are filed into their mapped structure subfolder by group,
    with a safe fallback to '05 Forms (Populated)/<group>' when unmapped."""
    root = tmp_path / "root"
    out = tmp_path / "out"
    build_sample_root(root)

    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    monkeypatch.setenv("OUTPUT_ROOT", str(out))
    monkeypatch.setenv("STAMP_DATE", "false")

    import app.config as config
    config.get_settings.cache_clear()
    config.load_departments.cache_clear()
    config.load_field_map.cache_clear()
    config.load_filing_map.cache_clear()

    from app.engine import generator, structure
    tree = structure.forms_tree()
    selected = [f["rel"] for grp in tree for f in grp["files"]]
    assert selected

    data = ProjectData(
        project_number="PJ-777777",
        site="Asda Leeds",
        department="M&E",
        structure="standard",
    )
    report = generator.generate(data, selected)
    dest = Path(report.output_dir)

    hs = dest / "04 - Health & Safety"
    phase = dest / "02 - Construction Phase"
    populated = dest / "05 Forms (Populated)"

    # OHS group -> Health & Safety (docx + the pptx Fire Plan)
    assert any("Site Safety" in f.name for f in hs.rglob("*") if f.is_file())
    assert any(f.suffix.lower() == ".pptx" for f in hs.rglob("*") if f.is_file())
    # QUAL/MS 9.4 (Project Directory) -> Construction Phase zone (per the Library Map)
    assert any("Project Notification" in f.name for f in phase.rglob("*") if f.is_file())
    # ENV + IMS have no mapped home -> fall back to the populated area
    assert any("Waste Note" in f.name for f in populated.rglob("*") if f.is_file())
    assert any("Weekly Inspection" in f.name for f in populated.rglob("*") if f.is_file())

    # Nothing errored and every filed form keeps the PJ prefix
    assert report.error_count == 0
    forms = [f for f in dest.rglob("*")
             if f.is_file() and f.suffix.lower() in (".docx", ".xlsx", ".pptx")]
    assert forms and all(f.name.startswith("PJ-777777 - ") for f in forms)


def test_filing_map_fallback_when_folder_absent(tmp_path, monkeypatch):
    """If a mapped destination folder doesn't exist in the structure, the form
    still lands safely in the populated fallback (never lost)."""
    root = tmp_path / "root"
    out = tmp_path / "out"
    build_sample_root(root)
    # Commercial has no '04 - Health & Safety' in the fixture, so OHS must fall back.
    (root / "00 Master Projects FIle Template 0524 (Opt 2)" / "Commercial" / "01 - Order").mkdir(parents=True, exist_ok=True)

    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    monkeypatch.setenv("OUTPUT_ROOT", str(out))
    import app.config as config
    config.get_settings.cache_clear()
    config.load_departments.cache_clear()
    config.load_field_map.cache_clear()
    config.load_filing_map.cache_clear()

    from app.engine import generator, structure
    tree = structure.forms_tree()
    ohs = [f["rel"] for grp in tree if grp["code"] == "OHS" for f in grp["files"]]
    assert ohs
    data = ProjectData(project_number="PJ-888888", site="Test Site",
                       department="Commercial", structure="standard")
    report = generator.generate(data, ohs)
    dest = Path(report.output_dir)
    populated = dest / "05 Forms (Populated)"
    assert any(populated.rglob("*.docx")), "OHS should fall back when no H&S folder exists"
    assert report.error_count == 0


# --------------------------------------------------------------------------- #
# Library-map (per-form) filing for the Construction structure.
# --------------------------------------------------------------------------- #
def _build_construction_root(root: Path) -> None:
    """A Construction TEMPLATES_ROOT with the exact subfolders the sample forms map to."""
    from tests.fixtures import (
        make_lookup_workbooks, make_docx_form, make_docx_simple,
        make_xlsx_form, make_pptx_form,
    )
    make_lookup_workbooks(root)
    master = root / "00 Master Projects FIle Template 0524 (Opt 2)" / "Construction"
    for sub in (
        "02 - Construction Phase/04 - Project Directory",
        "04 - Health & Safety/04 - Construction Phase Plan",
        "04 - Health & Safety/05 - Fire",
        "04 - Health & Safety/13 - Inspections",
        "04 - Health & Safety/19 - Waste",
    ):
        (master / sub).mkdir(parents=True, exist_ok=True)
    forms = root / "05 Forms"
    make_docx_form(forms / "01 OHS" / "HS 11.2 - Site Safety & Env Mgt Plan.docx")
    make_xlsx_form(forms / "02 ENV" / "ENV 1.2 - Waste Note.xlsx")
    make_docx_simple(forms / "03 QUAL" / "MS 9.4 Project Notification.docx")
    make_pptx_form(forms / "01 OHS" / "HS 13.1 - Site Fire Plan.pptx")
    make_docx_form(forms / "04 IMS" / "IMS 4.3 - PM Weekly Inspection.docx")


def _clear_config_caches():
    import app.config as config
    for fn in (config.get_settings, config.load_departments, config.load_field_map,
               config.load_filing_map, config.load_library_map):
        fn.cache_clear()


def test_library_map_files_construction_to_exact_subfolders(tmp_path, monkeypatch):
    """Each Construction form lands in its exact Library-Map subfolder, and the whole
    pack is bundled into a single downloadable zip with every filename PJ-prefixed."""
    import zipfile
    root = tmp_path / "root"
    out = tmp_path / "out"
    _build_construction_root(root)

    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    monkeypatch.setenv("OUTPUT_ROOT", str(out))
    monkeypatch.setenv("STAMP_DATE", "false")
    _clear_config_caches()

    from app.engine import generator, structure
    tree = structure.forms_tree()
    selected = [f["rel"] for grp in tree for f in grp["files"]]
    assert selected

    data = ProjectData(project_number="PJ-654321", site="Asda Stockton",
                       department="Construction", structure="standard")
    report = generator.generate(data, selected)
    dest = Path(report.output_dir)

    def filed(name_sub, relfolder):
        folder = dest / relfolder
        return folder.is_dir() and any(
            name_sub in f.name for f in folder.rglob("*") if f.is_file()
        )

    assert filed("Site Safety", "04 - Health & Safety/04 - Construction Phase Plan")   # HS 11.2
    assert filed("Waste Note", "04 - Health & Safety/19 - Waste")                       # ENV 1.2
    assert filed("Project Notification", "02 - Construction Phase/04 - Project Directory")  # MS 9.4
    assert filed("Site Fire Plan", "04 - Health & Safety/05 - Fire")                    # HS 13.1 (pptx)
    assert filed("Weekly Inspection", "04 - Health & Safety/13 - Inspections")          # IMS 4.3 (prefix)

    # Everything resolved -> nothing dumped into the populated fallback.
    populated = dest / "05 Forms (Populated)"
    assert not (populated.exists() and any(populated.rglob("*")))

    # Single zip, contains the filed forms in their subfolders, no errors, all prefixed.
    assert report.error_count == 0
    assert Path(report.zip_path).is_file()
    with zipfile.ZipFile(report.zip_path) as zf:
        names = zf.namelist()
    assert any("04 - Construction Phase Plan" in n and "Site Safety" in n for n in names)
    forms = [f for f in dest.rglob("*")
             if f.is_file() and f.suffix.lower() in (".docx", ".xlsx", ".pptx")]
    assert forms and all(f.name.startswith("PJ-654321 - ") for f in forms)


def test_library_map_fallback_when_subfolder_absent(tmp_path, monkeypatch):
    """When a mapped subfolder (and the coarse group folder) is absent, the form still
    lands safely in the populated fallback so nothing is ever lost."""
    from tests.fixtures import make_lookup_workbooks, make_docx_form
    root = tmp_path / "root"
    out = tmp_path / "out"
    make_lookup_workbooks(root)
    # Construction master WITHOUT '04 - Health & Safety' (and its Fire subfolder).
    master = root / "00 Master Projects FIle Template 0524 (Opt 2)" / "Construction"
    (master / "01 - Pre Construction").mkdir(parents=True, exist_ok=True)
    make_docx_form(root / "05 Forms" / "01 OHS" / "HS 3.1 - Fire Risk Assessment.docx")

    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    monkeypatch.setenv("OUTPUT_ROOT", str(out))
    _clear_config_caches()

    from app.engine import generator, structure
    tree = structure.forms_tree()
    selected = [f["rel"] for grp in tree for f in grp["files"]]
    assert selected

    data = ProjectData(project_number="PJ-999000", site="Test Site",
                       department="Construction", structure="standard")
    report = generator.generate(data, selected)
    dest = Path(report.output_dir)
    populated = dest / "05 Forms (Populated)"
    assert any(populated.rglob("*.docx")), "form should fall back when its subfolder is absent"
    assert report.error_count == 0


def test_library_map_root_zone_health_and_safety_department(tmp_path, monkeypatch):
    """The 'Health & Safety' department keeps its H&S folders at the top level (no
    '04 - Health & Safety/' prefix); a form must file into that root subfolder."""
    from tests.fixtures import make_lookup_workbooks, make_pptx_form
    root = tmp_path / "root"
    out = tmp_path / "out"
    make_lookup_workbooks(root)
    master = root / "00 Master Projects FIle Template 0524 (Opt 2)" / "Health & Safety"
    (master / "05 - Fire").mkdir(parents=True, exist_ok=True)
    make_pptx_form(root / "05 Forms" / "01 OHS" / "HS 13.1 - Site Fire Plan.pptx")  # HS 13 -> Fire

    monkeypatch.setenv("TEMPLATES_ROOT", str(root))
    monkeypatch.setenv("OUTPUT_ROOT", str(out))
    _clear_config_caches()

    from app.engine import generator, structure
    tree = structure.forms_tree()
    selected = [f["rel"] for grp in tree for f in grp["files"]]
    data = ProjectData(project_number="PJ-777001", site="HS Only",
                       department="Health & Safety", structure="standard")
    report = generator.generate(data, selected)
    dest = Path(report.output_dir)

    assert (dest / "05 - Fire").is_dir()
    assert any("Site Fire Plan" in f.name for f in (dest / "05 - Fire").rglob("*") if f.is_file())
    # Root zone: NOT nested under a '04 - Health & Safety' parent.
    assert not (dest / "04 - Health & Safety").exists()
    assert report.error_count == 0
