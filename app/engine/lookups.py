"""Dropdown data loaded from the CBES lookup workbooks (PMs / QS / Customers).

Files live in TEMPLATES_ROOT and are configured in config/departments.json under
"lookups". Missing files degrade gracefully to an empty list so the form still
works (the fields fall back to free text).
"""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path

import openpyxl

from ..config import get_settings, load_departments


def _read_column(path: Path, sheet: str, name_col: str) -> list[str]:
    if not path.is_file():
        return []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows, [])]
    try:
        idx = header.index(name_col)
    except ValueError:
        idx = 0
    out: list[str] = []
    for row in rows:
        if idx < len(row) and row[idx] not in (None, ""):
            out.append(str(row[idx]).strip())
    # De-duplicate, keep order, sort case-insensitively for a tidy dropdown.
    seen, uniq = set(), []
    for v in out:
        if v.lower() not in seen:
            seen.add(v.lower())
            uniq.append(v)
    return sorted(uniq, key=str.lower)


@lru_cache
def _load(kind: str) -> tuple[str, ...]:
    cfg = load_departments().get("lookups", {}).get(kind)
    if not cfg:
        return tuple()
    path = get_settings().TEMPLATES_ROOT / cfg["file"]
    return tuple(_read_column(path, cfg.get("sheet", ""), cfg.get("name_col", "")))


def project_managers() -> list[str]:
    return list(_load("project_managers"))


def quantity_surveyors() -> list[str]:
    return list(_load("quantity_surveyors"))


def customers() -> list[str]:
    return list(_load("customers"))


def clear_cache() -> None:
    _load.cache_clear()
