"""Stamp project fields into a .pptx file (tables + inline 'Label: value' text)."""
from __future__ import annotations

from pptx import Presentation

from .fields import FieldMatcher, inline_replacement


def _stamp_table(table, matcher: FieldMatcher) -> int:
    written = 0
    for row in table.rows:
        cells = list(row.cells)
        for i, cell in enumerate(cells):
            val = matcher.value_for(cell.text)
            if val is None or i + 1 >= len(cells):
                continue
            target = cells[i + 1]
            if matcher.is_label(target.text):
                continue
            target.text = val
            written += 1
    return written


def _stamp_text_frame(tf, matcher: FieldMatcher) -> int:
    written = 0
    for para in tf.paragraphs:
        text = "".join(run.text for run in para.runs)
        if not text:
            continue
        new = inline_replacement(text, matcher)
        if new is not None and new != text:
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            written += 1
    return written


def stamp_pptx(path: str, matcher: FieldMatcher) -> int:
    prs = Presentation(path)
    written = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                written += _stamp_table(shape.table, matcher)
            elif shape.has_text_frame:
                written += _stamp_text_frame(shape.text_frame, matcher)
    if written:
        prs.save(path)
    return written
