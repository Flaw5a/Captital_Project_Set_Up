"""Stamp project fields into a .docx file (tables + headers/footers + inline)."""
from __future__ import annotations

from docx import Document
from docx.table import Table

from .fields import FieldMatcher, inline_replacement


def _set_cell_text(cell, value: str) -> None:
    """Write `value` into a table cell while keeping the first run's formatting."""
    paras = cell.paragraphs
    if not paras:
        cell.text = value
        return
    p = paras[0]
    if p.runs:
        p.runs[0].text = value
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(value)
    # Clear any additional paragraphs in the cell (stale multi-line values).
    for extra in paras[1:]:
        for r in extra.runs:
            r.text = ""


def _stamp_table(table: Table, matcher: FieldMatcher) -> int:
    written = 0
    for row in table.rows:
        cells = row.cells
        seen: set = set()
        deduped = []
        for c in cells:
            key = id(c._tc)
            if key in seen:
                continue
            seen.add(key)
            deduped.append(c)
        for i, cell in enumerate(deduped):
            val = matcher.value_for(cell.text)
            if val is None:
                continue
            if i + 1 >= len(deduped):
                continue
            target = deduped[i + 1]
            # Don't overwrite a cell that is itself a known label.
            if matcher.is_label(target.text):
                continue
            _set_cell_text(target, val)
            written += 1
        # Recurse into nested tables.
        for cell in deduped:
            for nt in cell.tables:
                written += _stamp_table(nt, matcher)
    return written


def _stamp_paragraphs(paragraphs, matcher: FieldMatcher) -> int:
    written = 0
    for p in paragraphs:
        text = p.text
        new = inline_replacement(text, matcher)
        if new is not None and new != text:
            if p.runs:
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new)
            written += 1
    return written


def stamp_docx(path: str, matcher: FieldMatcher) -> int:
    """Open, stamp and save in place. Returns number of fields written."""
    doc = Document(path)
    written = 0
    for table in doc.tables:
        written += _stamp_table(table, matcher)
    written += _stamp_paragraphs(doc.paragraphs, matcher)
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            for table in hdr.tables:
                written += _stamp_table(table, matcher)
            written += _stamp_paragraphs(hdr.paragraphs, matcher)
        for ftr in (section.footer, section.first_page_footer, section.even_page_footer):
            for table in ftr.tables:
                written += _stamp_table(table, matcher)
            written += _stamp_paragraphs(ftr.paragraphs, matcher)
    if written:
        doc.save(path)
    return written
