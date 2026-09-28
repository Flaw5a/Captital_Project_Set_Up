"""Stamp project fields into a .xlsx file.

For each worksheet we scan used cells; when a cell's text matches a label we write
the value into the cell immediately to the right (unless that cell is itself a
label, in which case we try the cell below). Merged cells are handled by writing
to the top-left anchor of the target merge range.
"""
from __future__ import annotations

import openpyxl

from .fields import FieldMatcher


def _merged_anchor(ws, row: int, col: int):
    """Return (row, col) of the anchor cell if (row, col) is inside a merged range."""
    coord = ws.cell(row=row, column=col).coordinate
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return rng.min_row, rng.min_col
    return row, col


def _write(ws, row: int, col: int, value: str) -> bool:
    r, c = _merged_anchor(ws, row, col)
    cell = ws.cell(row=r, column=c)
    cell.value = value
    return True


def stamp_xlsx(path: str, matcher: FieldMatcher) -> int:
    wb = openpyxl.load_workbook(path)  # keep_vba handled separately for .xlsm
    written = 0
    for ws in wb.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                if not isinstance(cell.value, str):
                    continue
                val = matcher.value_for(cell.value)
                if val is None:
                    continue
                # Prefer the cell to the right.
                if col + 1 <= max_col + 1:
                    right = ws.cell(row=row, column=col + 1)
                    if not matcher.is_label(str(right.value or "")):
                        if _write(ws, row, col + 1, val):
                            written += 1
                            continue
                # Fall back to the cell below.
                below = ws.cell(row=row + 1, column=col)
                if not matcher.is_label(str(below.value or "")):
                    if _write(ws, row + 1, col, val):
                        written += 1
    if written:
        wb.save(path)
    return written
